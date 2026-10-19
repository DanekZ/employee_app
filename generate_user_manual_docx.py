import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_callout(doc, text, title="INFORMASI PENTING", bg_color="F9FAFB", border_color="800020"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Set left border color and width
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.bold = True
    run_t.font.name = 'Arial'
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(128, 0, 32) # Maroon #800020
    
    run_body = p.add_run(text)
    run_body.font.name = 'Arial'
    run_body.font.size = Pt(9.5)
    run_body.font.color.rgb = RGBColor(55, 65, 81)
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(6)

def build_user_manual():
    doc = Document()

    # Page Setup (Margining 1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(20)
    title_p.paragraph_format.space_after = Pt(6)
    
    run_title = title_p.add_run("PANDUAN PENGGUNAAN APLIKASI KARYAWAN")
    run_title.bold = True
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGBColor(128, 0, 32) # Brand Maroon #800020

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(20)
    run_sub = sub_p.add_run("Sistem Manajemen Absensi GPS, Pengajuan Izin, Lembur, Dinas Luar, Approval & Penggajian (Payroll)\nPT Sinarta MJS — Edisi Pembaruan Terkini (Oktober 2026)")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(11.5)
    run_sub.font.color.rgb = RGBColor(163, 8, 47) # Brand Crimson #A3082F

    # Divider line
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(16)
    run_div = p_div.add_run("_________________________________________________________________________________")
    run_div.font.color.rgb = RGBColor(229, 231, 235)

    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(128, 0, 32)
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(12.5)
        run.font.color.rgb = RGBColor(163, 8, 47)
        return h

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.bold = True
            run_b.font.name = 'Arial'
            run_b.font.size = Pt(10)
            run_b.font.color.rgb = RGBColor(31, 41, 55)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(55, 65, 81)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.bold = True
            run_b.font.name = 'Arial'
            run_b.font.size = Pt(10)
            run_b.font.color.rgb = RGBColor(31, 41, 55)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(55, 65, 81)
        return p

    # --- SECTION 1: PENDAHULUAN ---
    add_heading_1("1. PENDAHULUAN")
    add_body("Buku panduan ini merupakan acuan operasional resmi dalam menggunakan Aplikasi Karyawan Sinarta MJS. Aplikasi ini mencakup seluruh ekosistem kerja SDM, mulai dari pencatatan kehadiran presensi online berbasis koordinat geolokasi (GPS), permohonan izin/cuti, pengajuan lembur, penugasan perjalanan dinas, modul persetujuan (approval) atasan, hingga pengelolaan penggajian (payroll) dan pencetakan slip gaji digital.")

    add_heading_2("1.1 Peran dan Hak Akses Pengguna")
    
    table_role = doc.add_table(rows=4, cols=3)
    table_role.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table_role.rows[0].cells
    headers = ["Peran (Role)", "Pengguna Sasaran", "Hak Akses & Cakupan Fitur"]
    
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        set_cell_background(hdr_cells[i], "800020")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        for run in p.runs:
            run.font.bold = True
            run.font.name = 'Arial'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(255, 255, 255)

    roles_data = [
        ("Karyawan", "Seluruh Staf Operasional / Karyawan", 
         "• Presensi Check-in & Check-out berbasis koordinat GPS & tautan Google Maps.\n"
         "• Pengajuan Izin/Cuti, Lembur, dan Perjalanan Dinas Luar Kantor.\n"
         "• Memantau status pengajuan pribadi (Pending, Disetujui, Ditolak).\n"
         "• Mengakses menu 'Slip Gaji Saya', melihat rincian gaji, dan mencetak slip gaji resmi."),
        ("Atasan", "Manajer, Supervisor, & Pimpinan Divisi", 
         "• Seluruh akses Karyawan.\n"
         "• Dashboard Approval: Menyetujui (Approve) atau Menolak (Reject) pengajuan Izin, Lembur, dan Dinas bawahan.\n"
         "• Rekap Presensi Tim: Memantau kehadiran, jam kerja, dan titik lokasi seluruh tim bawahan.\n"
         "• Laporan Visualisasi Presensi: Statistik disiplin, keterlambatan, dan performa tim.\n"
         "• Modul Manajemen Penggajian (Payroll): Konfigurasi tarif global, input gaji pokok karyawan, generate massal slip gaji bulanan, dan publikasi slip gaji."),
        ("Admin", "Administrator Sistem & Divisi HRD", 
         "• Akses pengelolaan absensi dan pengajuan tim secara menyeluruh.\n"
         "• Akses ke halaman approval, rekap absensi, dan laporan visualisasi performa seluruh staf.\n"
         "• Terdaftar dalam sistem presensi dan laporan statistik tim.")
    ]

    for row_idx, data in enumerate(roles_data, start=1):
        row_cells = table_role.rows[row_idx].cells
        bg_c = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(data):
            row_cells[col_idx].text = cell_value
            set_cell_background(row_cells[col_idx], bg_c)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[col_idx].paragraphs[0]
            for run in p.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(55, 65, 81)

    add_body("")
    add_heading_2("1.2 Kebutuhan Perangkat & Izin Akses GPS")
    add_bullet(" Smartphone (Android / iOS) atau Laptop / Komputer dengan koneksi internet aktif.", "Perangkat: ")
    add_bullet(" Google Chrome, Mozilla Firefox, Microsoft Edge, atau Safari versi terbaru.", "Peramban (Browser): ")
    add_bullet(" Pengguna wajib mengizinkan peramban mengakses lokasi (Location Permission / GPS) saat tombol Check-in ditekan.", "Izin Lokasi (Wajib): ")

    create_callout(doc, "Saat pertama kali mengakses tombol 'Absen Masuk' atau 'Absen Keluar', peramban web akan menampilkan jendela izin lokasi: pilih 'Allow' atau 'Izinkan Saat Menggunakan'. Tanpa izin GPS, presensi tidak dapat direkam oleh sistem keamanan.", title="PENTING: IZIN GEOLOKASI (GPS)")

    # --- SECTION 2: PANDUAN MODUL KARYAWAN ---
    add_heading_1("2. PANDUAN PENGGUNA — MODUL KARYAWAN")

    add_heading_2("2.1 Autentikasi (Login & Logout) & Alur Masuk Otomatis")
    add_bullet(" Buka peramban web dan kunjungi tautan aplikasi karyawan.", "1. Buka Aplikasi: ")
    add_bullet(" Masukkan Email terdaftar dan Kata Sandi (Password) Anda, lalu klik tombol Masuk.", "2. Masukkan Kredensial: ")
    add_bullet(" Sistem cerdas secara otomatis mengarahkan Karyawan ke halaman Absensi Mandiri, dan mengarahkan Atasan langsung ke Dashboard Approval.", "3. Redirect Otomatis: ")
    add_bullet(" Untuk mengakhiri sesi kerja, klik nama profil di bilah navigasi kanan atas lalu pilih 'Keluar' (Log Out).", "4. Keluar Sistem: ")

    add_heading_2("2.2 Melakukan Presensi Harian (Check-in & Check-out)")
    add_bullet(" Masuk ke menu 'Absensi' pada navigasi utama.", "Langkah 1 (Buka Menu): ")
    add_bullet(" Untuk memulai jam kerja harian, klik tombol 'Absen Masuk'. Pastikan posisi GPS terdeteksi.", "Langkah 2 (Absen Masuk): ")
    add_bullet(" Sistem memvalidasi waktu masuk: kehadiran sebelum atau tepat pukul 08:00 WIB berstatus 'Hadir' (Hijau), sedangkan kehadiran setelah pukul 08:00 WIB berstatus 'Telat' (Kuning).", "Langkah 3 (Penentuan Status): ")
    add_bullet(" Pada akhir jam kerja kantor, klik tombol 'Absen Keluar' untuk mencatat jam pulang.", "Langkah 4 (Absen Keluar): ")
    add_bullet(" Pada bagian bawah halaman, periksa tabel 'Riwayat Bulan Ini' untuk meninjau jam kerja harian dan tautan titik peta Google Maps lokasi presensi Anda.", "Langkah 5 (Pemeriksaan Riwayat): ")

    add_heading_2("2.3 Pengajuan Izin / Cuti Online")
    add_body("Digunakan ketika berhalangan hadir atau membutuhkan izin meninggalkan kantor.")
    add_bullet(" Klik menu 'Pengajuan' -> pilih 'Izin'.", "1. Navigasi: ")
    add_bullet(" Klik tombol '+ Buat Pengajuan Izin'.", "2. Buka Formulir: ")
    add_bullet(" Tentukan Jenis Izin yang sesuai: 'Tidak Masuk Kerja', 'Datang Terlambat', 'Pulang Lebih Awal', atau 'Keluar Kantor Sementara'.", "3. Pilih Kategori: ")
    add_bullet(" Isi Tanggal Mulai, Tanggal Selesai (jika izin berhari-hari), durasi/jam, serta uraian alasan izin yang jelas.", "4. Rincian Izin: ")
    add_bullet(" Klik 'Kirim Pengajuan'. Berkas pengajuan akan terkirim dengan status 'Pending' untuk diverifikasi atasan langsung.", "5. Kirim: ")

    add_heading_2("2.4 Pengajuan Lembur")
    add_body("Wajib diajukan oleh staf yang melaksanakan tugas lembur di luar jam kerja efektif agar kompensasi lembur dapat terhitung di slip gaji.")
    add_bullet(" Klik menu 'Pengajuan' -> pilih 'Lembur', lalu tekan tombol '+ Ajukan Lembur'.", "1. Navigasi & Form: ")
    add_bullet(" Tentukan Tanggal Lembur, Jam Mulai, Jam Selesai, Lokasi Lembur (misal: Kantor Pusat / Lapangan), dan uraian pekerjaan yang dikerjakan.", "2. Pengisian Form: ")
    add_bullet(" Klik 'Kirim Pengajuan Lembur'. Setelah disetujui oleh Atasan (Approved), total durasi jam lembur otomatis terakumulasi ke dalam perhitungan slip gaji bulanan Anda.", "3. Integrasi Penggajian: ")

    add_heading_2("2.5 Pengajuan Perjalanan Dinas Luar Kantor")
    add_body("Digunakan saat staf ditugaskan melakukan perjalanan dinas atau kunjungan luar kantor.")
    add_bullet(" Klik menu 'Pengajuan' -> pilih 'Dinas Luar' -> klik '+ Ajukan Dinas Luar'.", "1. Navigasi & Form: ")
    add_bullet(" Tentukan Tanggal Dinas, Alamat/Tujuan Dinas, Jam Keluar, Jam Kembali, Alat Transportasi (Kendaraan Dinas, Pribadi, atau Umum), serta maksud penugasan.", "2. Kelengkapan Data: ")
    add_bullet(" Kirim pengajuan. Pengajuan dinas yang disetujui (Approved) akan otomatis diperhitungkan sebagai kehadiran kerja sah dan menambah tunjangan dinas pada slip gaji.", "3. Status & Tunjangan: ")

    add_heading_2("2.6 Mengakses & Mencetak 'Slip Gaji Saya'")
    add_body("Setiap staf dapat memeriksa dokumen penerimaan gaji resmi secara transparan setelah slip gaji diterbitkan oleh Atasan/Manajemen.")
    add_bullet(" Klik menu 'Slip Gaji Saya' pada navigasi utama.", "1. Akses Menu: ")
    add_bullet(" Halaman menampilkan kartu slip gaji per periode (misal: 2026-09, 2026-10) dengan status 'Terbit'.", "2. Kartu Periode: ")
    add_bullet(" Klik tombol 'Lihat Detail' untuk membuka modal rincian slip gaji, yang menampilkan: Gaji Pokok, Tambahan Lembur (berdasarkan jam), Tunjangan Dinas, Potongan Alpha/Izin, dan Total Gaji Bersih (Take Home Pay).", "3. Rincian Lengkap: ")
    add_bullet(" Klik tombol 'Cetak Slip' untuk mencetak fisik atau menyimpan dokumen slip gaji resmi ke format PDF via peramban.", "4. Cetak Resmi: ")

    # --- SECTION 3: PANDUAN ATASAN & ADMIN ---
    add_heading_1("3. PANDUAN PENGGUNA — MODUL ATASAN & ADMIN")

    add_heading_2("3.1 Dashboard Approval (Persetujuan Cepat)")
    add_body("Atasan langsung dan Administrator dapat mengelola seluruh pengajuan tim kerja bawahan pada satu pintu terintegrasi.")
    add_bullet(" Buka menu 'Approval' pada bilah navigasi.", "1. Akses Dashboard: ")
    add_bullet(" Terdapat 3 tab utama: 'Pengajuan Izin Pending', 'Pengajuan Lembur Pending', dan 'Pengajuan Dinas Pending'.", "2. Tab Kategori: ")
    add_bullet(" Tabel menampilkan nama pemohon, tanggal kegiatan, durasi/waktu, keterangan, serta tombol aksi cepat.", "3. Tinjauan Berkas: ")

    add_heading_2("3.2 Tata Cara Menyetujui (Approve) & Menolak (Reject)")
    add_bullet(" Klik tombol 'Setujui' (warna Hijau) untuk menyetujui pengajuan. Status pengajuan langsung berubah menjadi 'Approved', nama penyetujui dan waktu approval otomatis tercatat di sistem.", "1. Menyetujui: ")
    add_bullet(" Klik tombol 'Tolak' (warna Merah) jika permohonan tidak sesuai. Status pengajuan otomatis berubah menjadi 'Rejected'.", "2. Menolak: ")

    create_callout(doc, "Data lembur dan dinas luar kantor yang disetujui (Approved) akan otomatis diperhitungkan oleh PayslipService saat Atasan mengkalkulasi slip gaji karyawan pada akhir periode bulan.", title="INTEGRASI OTOMATIS KE SISTEM PENGGAJIAN")

    add_heading_2("3.3 Rekapitulasi Presensi Tim & Titik Lokasi Peta")
    add_bullet(" Akses menu 'Approval' -> pilih sub-menu 'Rekap Absensi'.", "1. Buka Rekap Absensi: ")
    add_bullet(" Manfaatkan filter pencarian canggih: berdasarkan Nama Karyawan (karyawan & admin), Tanggal Spesifik, Bulan Kerja, atau Status Kehadiran (Hadir, Telat, Alpha).", "2. Filter Pencarian: ")
    add_bullet(" Klik tombol 'Buka Peta' pada setiap baris absensi untuk melihat titik koordinat Google Maps tempat karyawan melakukan check-in/out.", "3. Verifikasi Geolocation: ")

    add_heading_2("3.4 Laporan Visualisasi & Statistik Performa Kehadiran")
    add_bullet(" Buka sub-menu 'Laporan' pada navigasi Approval.", "1. Buka Laporan: ")
    add_bullet(" Memuat grafik harian presensi tim: Tren Hadir Tepat Waktu vs Terlambat sepanjang rentang tanggal terpilih.", "2. Grafik Tren Kehadiran: ")
    add_bullet(" Kartu Metrik Tim: Menampilkan Total Kehadiran, Total Hari Izin, Total Lembur, dan Total Kunjungan Dinas.", "3. Indikator Kunci: ")
    add_bullet(" Tabel Statistik Karyawan: Menampilkan persentase ketepatan waktu, jumlah keterlambatan, dan predikat kedisiplinan kerja masing-masing staf (termasuk staf role Karyawan dan Admin).", "4. Tabel Performa Individu: ")

    add_heading_2("3.5 Modul Manajemen Penggajian (Payroll)")
    add_body("Modul khusus bagi Atasan untuk mengelola seluruh aspek kompensasi finansial tim secara akurat dan transparan.")
    add_bullet(" Masuk ke menu 'Penggajian'. Tentukan Periode Bulan (format: YYYY-MM) melalui filter periode di bagian atas.", "1. Pilih Periode: ")
    add_bullet(" Tab 'Konfigurasi Tarif': Atasan dapat mengatur tarif standar global perusahaan: Tarif Lembur per Jam, Tarif Dinas per Jam, Potongan Alpha per Hari, dan Potongan Izin per Hari. Klik 'Simpan Pengaturan'.", "2. Atur Tarif Global: ")
    add_bullet(" Tab 'Kelola Gaji Pokok': Menampilkan daftar seluruh karyawan beserta gaji pokok saat ini. Klik tombol pensil/edit untuk mengubah gaji pokok karyawan melalui jendela popup modal.", "3. Atur Gaji Pokok Karyawan: ")
    add_bullet(" Tombol 'Kalkulasi / Generate Slip': Klik tombol ini untuk memerintahkan sistem menghitung otomatis seluruh slip gaji karyawan pada periode tersebut. Formula otomatis menggabungkan: Gaji Pokok + (Jam Lembur × Rate) + (Jam Dinas × Rate) - (Hari Alpha × Potongan Alpha) - (Hari Izin × Potongan Izin).", "4. Generate Massal: ")
    add_bullet(" Tab 'Daftar Slip Gaji': Menampilkan seluruh slip gaji yang telah dihitung. Atasan dapat mengklik 'Rincian' untuk memeriksa detail kalkulasi tiap individu.", "5. Verifikasi Rincian: ")
    add_bullet(" Tombol 'Terbitkan Slip Gaji': Setelah seluruh angka terverifikasi, klik tombol 'Terbitkan Slip Gaji'. Status slip gaji akan berubah menjadi 'Terbit' dan otomatis dapat dilihat serta dicetak oleh masing-masing karyawan di menu 'Slip Gaji Saya'.", "6. Publikasi ke Karyawan: ")

    # --- SECTION 4: FAQ & PEMECAHAN MASALAH ---
    add_heading_1("4. PERTANYAAN UMUM & SOLUSI MASALAH (FAQ)")

    table_faq = doc.add_table(rows=5, cols=2)
    table_faq.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_faq = table_faq.rows[0].cells
    hdr_faq[0].text = "Permasalahan / Kendala yang Sering Dihadapi"
    hdr_faq[1].text = "Solusi & Tindakan Penanganan"
    
    set_cell_background(hdr_faq[0], "800020")
    set_cell_background(hdr_faq[1], "800020")
    set_cell_margins(hdr_faq[0], top=100, bottom=100, left=120, right=120)
    set_cell_margins(hdr_faq[1], top=100, bottom=100, left=120, right=120)
    
    for c in hdr_faq:
        p = c.paragraphs[0]
        for run in p.runs:
            run.font.bold = True
            run.font.name = 'Arial'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(255, 255, 255)

    faqs = [
        ("Gagal melakukan absensi: 'Gagal mengambil lokasi / Izin lokasi ditolak'.",
         "1. Pastikan fitur GPS di HP atau laptop Anda sudah dinyalakan.\n"
         "2. Masuk ke Pengaturan Browser -> Privacy & Security -> Site Settings -> Location -> pilih 'Allow'.\n"
         "3. Muat ulang (Refresh) halaman web dan tekan tombol Absen kembali."),
        
        ("Nominal lembur atau dinas luar tidak muncul di slip gaji bulanan.",
         "Pastikan pengajuan lembur atau dinas luar Anda telah disetujui (Approved) oleh Atasan. Pengajuan yang masih berstatus 'Pending' atau 'Rejected' tidak akan dikalkulasikan ke dalam slip gaji."),
        
        ("Karyawan tidak dapat melihat menu 'Slip Gaji Saya' atau slip periode aktif belum ada.",
         "Slip gaji bulanan hanya akan tampil apabila Atasan telah melakukan proses 'Kalkulasi' dan menekan tombol 'Terbitkan Slip Gaji' pada periode bulan bersangkutan."),
        
        ("Lupa kata sandi (password) akun login.",
         "Silakan hubungi Administrator IT atau Divisi HRD untuk melakukan pengaturan ulang (reset) kata sandi akun Anda ke password standar perusahaan.")
    ]

    for idx, (q, a) in enumerate(faqs, start=1):
        r_cells = table_faq.rows[idx].cells
        bg_c = "F9FAFB" if idx % 2 == 1 else "FFFFFF"
        
        r_cells[0].text = q
        r_cells[1].text = a
        
        set_cell_background(r_cells[0], bg_c)
        set_cell_background(r_cells[1], bg_c)
        set_cell_margins(r_cells[0], top=80, bottom=80, left=120, right=120)
        set_cell_margins(r_cells[1], top=80, bottom=80, left=120, right=120)
        
        for cell in r_cells:
            p = cell.paragraphs[0]
            for run in p.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(55, 65, 81)

    # --- SECTION 5: KONTAK BANTUAN ---
    add_heading_1("5. INFORMASI KONTAK LAYANAN BANTUAN")
    add_body("Apabila Anda memerlukan panduan lebih lanjut, kendala akses akun, atau perbaikan data kepegawaian, silakan menghubungi layanan resmi berikut:")
    add_bullet(" hrd@sinarta-mjs.co.id", "Divisi HRD & Kepegawaian: ")
    add_bullet(" it-support@sinarta-mjs.co.id", "Tim Teknis & IT Support: ")
    add_bullet(" Senin s/d Jumat, Pukul 08:00 – 17:00 WIB", "Jam Layanan Operasional: ")

    output_doc = "Panduan_Penggunaan_Aplikasi_Karyawan.docx"
    doc.save(output_doc)
    print(f"User manual berhasil diperbarui: {output_doc}")

if __name__ == '__main__':
    build_user_manual()
